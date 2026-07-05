#!/usr/bin/env python3
"""Generate INSTRUCTIONS.docx for the fishing/craft content helper handoff."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "INSTRUCTIONS.docx"


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_bullet(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        p.add_run(text)
    else:
        p.add_run(text)


def main() -> None:
    doc = Document()

    # Title
    title = doc.add_heading("Into the Pond — Fishing & Craft Content Helper Brief", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.add_paragraph(
        "This package is your assignment brief for Into the Pond v3. "
        "You will produce content and design tables that engineering will implement directly. "
        "Work only inside the deliverables/ folder unless asked otherwise."
    )

    add_heading(doc, "1. Before you start", 1)
    add_bullet(doc, "Read every file listed in reference/REFERENCE_FILES.md (in order).", "Reference: ")
    add_bullet(doc, "Read reference/CURRENT_ENGINE_SUMMARY.md — note what is live today vs target design.", "Engine summary: ")
    add_bullet(doc, "Do not edit files under shared/, functions/, or src/ — deliver completed work in deliverables/.", "Do not: ")

    add_heading(doc, "2. Folder structure", 1)
    doc.add_paragraph("You received (or cloned) this folder:")
    structure = doc.add_paragraph()
    structure.add_run(
        "docs/handoff/fishing-craft-content-helper/\n"
        "  INSTRUCTIONS.docx          ← this document\n"
        "  reference/\n"
        "    REFERENCE_FILES.md       ← files to read in the codebase\n"
        "    CURRENT_ENGINE_SUMMARY.md\n"
        "  deliverables/              ← fill in and return these files\n"
        "    01-audio-manifest.csv\n"
        "    02-fishing-outcomes-matrix.csv\n"
        "    03-craft-journey-scenarios.md\n"
        "    04-fishing-outcome-messages.md\n"
        "    05-probability-tables.md\n"
    ).font.name = "Courier New"
    structure.runs[0].font.size = Pt(9)

    add_heading(doc, "3. Your five tasks", 1)

    # Task 1
    add_heading(doc, "Task 1 — Audio prep", 2)
    doc.add_paragraph(
        "There is no fishing audio in the app yet. Source or produce short, warm, non-urgent clips "
        "(field-journal tone — soft, natural, not arcade SFX)."
    )
    add_bullet(doc, "Fill in deliverables/01-audio-manifest.csv — one row per clip.")
    add_bullet(doc, "Deliver final audio files separately (MP3 or M4A), named exactly as in the filename column.")
    add_bullet(doc, "Suggested moments: cast splash, waiting ambient (2 min loop), claim catch/duplicate/miss, craft begin/ready/equip.")
    add_bullet(doc, "Use assets/audio/journal/ as tone reference only.")
    p = doc.add_paragraph()
    p.add_run("Cast timer in app: ").bold = True
    p.add_run("2 minutes between cast and claim.")

    # Task 2
    add_heading(doc, "Task 2 — Fishing outcomes spreadsheet (150 rows)", 2)
    doc.add_paragraph(
        "Complete deliverables/02-fishing-outcomes-matrix.csv. Rows are pre-filled with "
        "10 rods × 3 baits × 5 Wonder levels (0, 25, 45, 65, 90). You fill the Expected columns and Notes."
    )
    add_bullet(doc, "expected_creature_pool — describe pool after wonder gate (size, tier) or write EMPTY if gated.")
    add_bullet(doc, "expected_catch — Yes / No / Probabilistic (include % when pool is open).")
    add_bullet(doc, "expected_miss_reason — wonder_gate, chance, or n/a.")
    add_bullet(doc, "notes — edge cases (e.g. wildcard @ 0W + premium bait; locked rod; duplicate tab).")
    doc.add_paragraph(
        "Important: Today all three baits behave the same in code. Mark rows CURRENT vs TARGET where "
        "your answers assume the new bait rules in Task 5."
    )

    # Task 3
    add_heading(doc, "Task 3 — Craft progression test scenarios (8 journeys)", 2)
    doc.add_paragraph(
        "Complete deliverables/03-craft-journey-scenarios.md. Write eight realistic parent timelines "
        "with Firestore JSON snapshots at key days."
    )
    add_bullet(doc, "Journey 1 — Steady Module 1 finisher (craft rare_fire)")
    add_bullet(doc, "Journey 2 — Abandons mid-craft (never collects)")
    add_bullet(doc, "Journey 3 — Collects very late (timer done days ago)")
    add_bullet(doc, "Journey 4 — Wonder gate fisher (owned rod, low currentWonder)")
    add_bullet(doc, "Journey 5 — Cast with locked rod (server rejection)")
    add_bullet(doc, "Journey 6 — Premium migrator (backfill from completedLessons)")
    add_bullet(doc, "Journey 7 — Wildcard journey gift (four rares → rare_wildcard ready)")
    add_bullet(doc, "Journey 8 — Epic post-curriculum (all lessons + epic_fire craft)")
    doc.add_paragraph(
        "Show parts and Wonder at each step. Parts: +1 per lesson, +5 cluster bonus. "
        "Craft costs: see shared/sanctuary/progression/craftCosts.ts."
    )

    # Task 4
    add_heading(doc, "Task 4 — Fishing outcome messages", 2)
    doc.add_paragraph(
        "Complete deliverables/04-fishing-outcome-messages.md."
    )
    add_bullet(doc, "Field-journal voice, past tense, max 2 sentences.")
    add_bullet(doc, "5 variants each: new catch, duplicate, miss (wonder gate), miss (chance). Use {creatureName} where needed.")
    add_bullet(doc, "10 first-cast messages — one per domain rod ID, referencing element + module theme.")

    # Task 5
    add_heading(doc, "Task 5 — Probability tables", 2)
    doc.add_paragraph(
        "Complete deliverables/05-probability-tables.md (and attach extra CSVs if needed). "
        "This is the implementation spec — engineering will build from your tables."
    )
    add_bullet(doc, "Basic rod: ~30 common creatures, roughly equal weight.")
    add_bullet(doc, "Rare element: ~14 creatures in 6/5/3 sub-tiers; target ~1/15 top-rare, ~1/6 mid-rare.")
    add_bullet(doc, "Epic: shifted weights — top-rare ~1/6, mid-rare ~1/3.")
    add_bullet(doc, "Wildcard: all elements at lower per-creature weight than dedicated rods.")
    add_bullet(doc, "Pity: 10 consecutive chance misses → guaranteed catch; 20 epic casts without top-rare → +15% weight per cast.")
    add_bullet(doc, "Bait: left +5%, middle +12% + remove 1 gate tier, right +20% + remove 2 gate tiers.")
    doc.add_paragraph(
        "Include worked maths: Rod type × Creature sub-tier × Base weight × Bait modifier × Pity-adjusted weight."
    )

    add_heading(doc, "4. Rod catalog (quick reference)", 1)
    table = doc.add_table(rows=11, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Domain rod ID"
    hdr[1].text = "Tier"
    hdr[2].text = "Element"
    hdr[3].text = "Pool opens (currentWonder)"
    data = [
        ("basic", "basic", "any", "0"),
        ("rare_fire", "rare", "fire", "40"),
        ("rare_water", "rare", "water", "40"),
        ("rare_wind", "rare", "wind", "40"),
        ("rare_electric", "rare", "electric", "40"),
        ("rare_wildcard", "rare", "any", "65"),
        ("epic_fire", "epic", "fire", "90"),
        ("epic_water", "epic", "water", "90"),
        ("epic_wind", "epic", "wind", "90"),
        ("epic_electric", "epic", "electric", "90"),
    ]
    for i, row in enumerate(data, start=1):
        cells = table.rows[i].cells
        for j, val in enumerate(row):
            cells[j].text = val

    add_heading(doc, "5. How to submit", 1)
    doc.add_paragraph(
        "Return the completed deliverables/ folder (all five files plus any audio assets and extra CSVs). "
        "Keep filenames unchanged so we can diff your work."
    )
    add_bullet(doc, "01-audio-manifest.csv + audio files")
    add_bullet(doc, "02-fishing-outcomes-matrix.csv (150 rows filled)")
    add_bullet(doc, "03-craft-journey-scenarios.md (8 journeys)")
    add_bullet(doc, "04-fishing-outcome-messages.md")
    add_bullet(doc, "05-probability-tables.md")

    add_heading(doc, "6. Questions", 1)
    doc.add_paragraph(
        "If something is ambiguous, note your assumption in the notes column or in a short comment at the top "
        "of the relevant deliverable file. Prefer documenting assumptions over blocking."
    )

    doc.add_paragraph()
    footer = doc.add_paragraph("Into the Pond v3 — content helper brief")
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
